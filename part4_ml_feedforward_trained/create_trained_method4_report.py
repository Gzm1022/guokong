from pathlib import Path
import csv

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
PY_DIR = ROOT / "python"
FIG = PY_DIR / "figures"
OUT = ROOT / "方法4_随机森林软测量前馈串级PID仿真实验报告.docx"


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    for name, size, color in [
        ("Heading 1", 16, RGBColor(31, 78, 121)),
        ("Heading 2", 13, RGBColor(47, 84, 150)),
    ]:
        st = doc.styles[name]
        st.font.name = "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color


def p(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Pt(21)
    para.paragraph_format.line_spacing = 1.25
    run = para.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)
    return para


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    run = para.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_text(cell, text, bold=False):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(str(text))
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(9.5)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        shade(tbl.rows[0].cells[i], "D9EAF7")
        cell_text(tbl.rows[0].cells[i], h, True)
    for row in rows:
        cells = tbl.add_row().cells
        for i, item in enumerate(row):
            cell_text(cells[i], item)
    return tbl


def figure(doc, path, caption, width=6.3):
    if path.exists():
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(path), width=Inches(width))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(9)
        run.italic = True
    else:
        p(doc, f"图片缺失：{path}")


def read_metric_pairs(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return {row["metric"]: row["value"] for row in csv.DictReader(f)}


def read_control_metrics(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def fmt(row, key, nd=3):
    return f"{float(row[key]):.{nd}f}"


def main():
    model_metrics = read_metric_pairs(PY_DIR / "ml_model_metrics.csv")
    control_metrics = read_control_metrics(PY_DIR / "method4_trained_metrics.csv")

    name_map = {
        "Cascade PID": "串级 PID",
        "Conventional FF + Cascade PID": "普通前馈 + 串级 PID",
        "Random Forest FF + Cascade PID": "随机森林前馈 + 串级 PID",
    }
    control_rows = [
        [
            name_map[row["controller"]],
            fmt(row, "min_do"),
            fmt(row, "max_abs_error"),
            fmt(row, "recovery_time_s", 1),
            fmt(row, "iae_after_dist", 2),
            fmt(row, "air_command_energy", 2),
        ]
        for row in control_metrics
    ]

    doc = Document()
    style_doc(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("方法四训练版：随机森林软测量前馈 + 串级 PID\n污水曝气溶解氧控制仿真实验报告")
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("过程控制课程仿真实验 | 随机森林回归软测量 | 前馈-反馈复合控制")
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_heading("1 实验目的", level=1)
    p(doc, "本实验在普通前馈 + 串级 PID 的基础上，引入真正训练得到的机器学习软测量模型。由于课程设计缺少实际污水厂 SCADA 历史数据，本文基于前文建立的曝气 DO 机理模型构造多工况仿真历史样本，训练随机森林回归模型预测当前耗氧负荷，并将预测结果作为前馈信号接入串级 PID 控制系统。")
    p(doc, "该版本与直接手写前馈模型不同，包含训练数据生成、模型训练、测试集评价和控制系统验证四个环节，因此能够更完整地说明机器学习前馈在提升 DO 控制快速性方面的作用。")

    doc.add_heading("2 与前面方法的关系", level=1)
    p(doc, "方法四训练版仍沿用前面方法的基础被控对象：供气侧对象为一阶惯性环节，DO 主对象为一阶惯性加 30 s 纯滞后。控制目标仍为将 DO 稳定在 2.0 mg/L 附近。与方法三相比，方法四不直接依赖慢速 COD 在线测量，而是用随机森林模型根据实时易测变量预测耗氧负荷。")
    table(doc, ["项目", "方法三", "方法四训练版"], [
        ["前馈输入", "进水流量 Q、慢速污染物浓度 C", "Q、pH、电导率、温度、ORP、时间特征"],
        ["前馈模型", "人工设定耗氧量模型", "训练得到的随机森林回归模型"],
        ["控制结构", "普通前馈 + DO 外环 + 供气内环", "RF 软测量前馈 + DO 外环 + 供气内环"],
        ["主要目的", "提前补偿进水负荷扰动", "在污染物测量慢时更早估计负荷并补偿"],
    ])

    doc.add_heading("3 训练数据与模型", level=1)
    p(doc, "训练样本由仿真程序生成，共 8000 条。每条样本包含进水流量、pH、电导率、温度、ORP 以及时间周期特征，标签为由机理关系计算得到的真实耗氧负荷。样本中加入了随机工况变化和测量噪声，用于模拟污水厂 SCADA 数据的波动。")
    table(doc, ["输入变量", "物理含义"], [
        ["q_in", "归一化进水流量"],
        ["pH", "进水酸碱度，污染负荷升高时可能下降"],
        ["conductivity", "电导率，反映溶解性污染物变化"],
        ["temperature", "水温，影响微生物耗氧与传质"],
        ["orp", "氧化还原电位，反映水质氧化还原状态"],
        ["time_sin/time_cos", "时间周期特征，反映日周期运行规律"],
    ])
    table(doc, ["模型指标", "数值"], [
        ["训练样本数", model_metrics["train_rows"]],
        ["测试样本数", model_metrics["test_rows"]],
        ["MAE", f"{float(model_metrics['mae']):.4f}"],
        ["RMSE", f"{float(model_metrics['rmse']):.4f}"],
        ["R2", f"{float(model_metrics['r2']):.4f}"],
    ])
    p(doc, "随机森林模型的测试集 R2 接近 1，说明模型能够较好地从实时易测变量中恢复耗氧负荷信息。需要说明的是，该结果基于仿真构造数据，不应等同于真实污水厂数据上的泛化性能。")

    doc.add_heading("4 控制仿真设计", level=1)
    p(doc, "控制仿真沿用方法三的进水负荷阶跃工况：t = 1200 s 时，进水流量由 1.00 增至 1.20，真实污染物浓度由 1.00 增至 1.40。普通前馈使用慢速污染物浓度测量，随机森林前馈使用训练好的模型预测耗氧负荷。")
    table(doc, ["控制方法", "说明"], [
        ["串级 PID", "仅依靠 DO 外环和供气流量内环反馈调节"],
        ["普通前馈 + 串级 PID", "使用 Q 与慢速 C 测量值计算前馈补偿"],
        ["随机森林前馈 + 串级 PID", "使用 RF 预测负荷计算前馈补偿，并保留 PID 反馈修正"],
    ])

    doc.add_heading("5 仿真结果分析", level=1)
    figure(doc, FIG / "trained_fig1_rf_prediction.png", "图1 随机森林软测量负荷预测效果")
    p(doc, "图1表明，随机森林预测负荷能够在进水扰动发生后快速接近真实耗氧负荷，为前馈通道提供更及时的补偿依据。")
    figure(doc, FIG / "trained_fig2_do_response.png", "图2 三种控制方法的 DO 响应对比")
    p(doc, "图2显示，随机森林前馈 + 串级 PID 的 DO 下跌幅度最小，且恢复速度明显快于普通前馈和单纯串级 PID。")
    figure(doc, FIG / "trained_fig3_air_setpoint.png", "图3 目标供气量对比")
    p(doc, "图3说明随机森林前馈在扰动发生后更早提高目标供气量，使供气内环提前动作。")
    figure(doc, FIG / "trained_fig4_control_signal.png", "图4 鼓风机控制指令对比")
    p(doc, "图4显示随机森林前馈方案在扰动初期给出更积极的鼓风机控制指令，从而换取更快的 DO 恢复。")
    figure(doc, FIG / "trained_fig5_performance_bar.png", "图5 抗扰动性能指标相对对比")
    p(doc, "图5将各指标按串级 PID 归一化为 100%。随机森林前馈方案在最大偏差、恢复时间和 IAE 上均最低，说明其快速性和抗扰动性能更好。")

    doc.add_heading("6 定量指标", level=1)
    table(doc, ["控制方法", "最低 DO / mg/L", "最大偏差 / mg/L", "恢复时间 / s", "IAE", "控制能量"], control_rows)

    doc.add_heading("7 结论", level=1)
    p(doc, "训练版方法四不是简单手写一个前馈公式，而是通过仿真历史样本训练随机森林软测量模型，并将训练好的模型接入前馈 + 串级 PID 控制结构。仿真结果表明，随机森林前馈能够在污染物浓度测量滞后的情况下更快估计耗氧负荷，从而提前增加供气量。")
    p(doc, "与普通前馈 + 串级 PID 相比，随机森林前馈 + 串级 PID 进一步降低了 DO 最大偏差和 IAE，并显著缩短恢复时间。由于训练数据为仿真构造数据，真实工程应用仍需使用实际 SCADA 历史数据重新训练和验证模型。")

    doc.save(OUT)
    print(f"Report saved: {OUT}")


if __name__ == "__main__":
    main()
