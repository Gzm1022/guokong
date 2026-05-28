from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
PY_DIR = ROOT / "python"
PY_FIG = PY_DIR / "figures"
MAT_FIG = ROOT / "matlab" / "figures"
OUT = ROOT / "方法3_普通前馈串级PID仿真实验报告.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(9)
    r.italic = True


def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)


def add_picture_if_exists(doc, path, caption, width=6.4):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        add_caption(doc, caption)
    else:
        add_body_paragraph(doc, f"图片缺失：{path}")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], "D9EAF7")
        set_cell_text(hdr[i], h, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    return table


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    for name, size, color in [
        ("Heading 1", 16, RGBColor(31, 78, 121)),
        ("Heading 2", 13, RGBColor(47, 84, 150)),
        ("Heading 3", 11, RGBColor(31, 78, 121)),
    ]:
        st = styles[name]
        st.font.name = "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True


def main():
    doc = Document()
    style_doc(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("方法三：普通前馈 + 串级 PID\n污水曝气溶解氧控制仿真实验报告")
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("过程控制课程仿真实验 | 被控量：溶解氧 DO | 控制对象：污水处理好氧池曝气系统")
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_paragraph()
    doc.add_heading("1 实验目的", level=1)
    add_body_paragraph(doc, "本实验针对污水处理好氧池中的曝气溶解氧控制问题，建立普通前馈 + 串级 PID 控制仿真模型。控制目标是在进水流量和污染物浓度发生突变时，维持水体中溶解氧 DO 稳定在 2.0 mg/L 附近，同时避免过度曝气造成能耗增加。")
    add_body_paragraph(doc, "方法三在方法二串级 PID 的基础上增加前馈补偿通道。前馈通道根据进水流量和污染物浓度估算耗氧负荷，在 DO 明显下降之前提前提高目标供气量；DO 外环反馈则负责消除前馈模型误差和不可测扰动。")

    doc.add_heading("2 控制方案与结构", level=1)
    add_body_paragraph(doc, "方法三采用“普通前馈 + 串级 PID”结构。外环为 DO 控制器，输出供气量反馈修正量；前馈模型根据进水扰动计算基础供气量；内环为供气流量控制器，用于快速调节鼓风机变频指令或空气阀门开度。")
    add_bullet(doc, "前馈输入：进水流量 Q、污染物浓度 C。")
    add_bullet(doc, "前馈输出：估算耗氧负荷对应的前馈供气量 q_ff。")
    add_bullet(doc, "反馈外环：根据 DO 设定值与测量值的偏差进行微调。")
    add_bullet(doc, "反馈内环：根据目标供气量与实际供气量的偏差调节鼓风机或阀门。")

    add_picture_if_exists(
        doc,
        MAT_FIG / "do_feedforward_cascade_diagram.png",
        "图1 方法三普通前馈 + 串级 PID Simulink 风格结构框图",
        width=6.7,
    )

    doc.add_heading("3 仿真模型", level=1)
    add_body_paragraph(doc, "供气侧对象采用一阶惯性模型，表示鼓风机或阀门控制指令到实际供气流量的动态过程。DO 主对象采用一阶惯性加纯滞后模型，表示供气进入水体后经过传质、混合和检测延迟再影响 DO 的过程。")
    add_table(doc, ["对象", "传递函数或模型", "说明"], [
        ["供气侧对象 G2", "1 / (15s + 1)", "鼓风机/阀门指令到实际供气流量"],
        ["纯滞后", "exp(-30s)", "供气传质与 DO 测量延迟"],
        ["DO 主对象 G1", "4 / (180s + 1)", "供气流量到 DO 的主动态"],
        ["DO 设定值", "2.0 mg/L", "好氧池典型控制目标"],
    ])

    doc.add_heading("4 前馈扰动模型", level=1)
    add_body_paragraph(doc, "进水扰动在 t = 1200 s 时加入，进水流量由 1.00 增加到 1.20，污染物浓度由 1.00 增加到 1.40。真实耗氧负荷与前馈估算模型故意设置为不完全一致，用于体现普通前馈存在模型误差。")
    add_table(doc, ["变量", "初值", "阶跃后", "含义"], [
        ["Q", "1.00", "1.20", "归一化进水流量"],
        ["C", "1.00", "1.40", "归一化污染物浓度"],
        ["d_load", "0", "负向扰动", "耗氧增加导致 DO 下降"],
    ])
    add_body_paragraph(doc, "真实耗氧负荷模型为：d_load = -[1.20(Q-Q0) + 1.60(C-C0)]。前馈估算模型为：q_ff = q_base + [1.05(Q_meas-Q0) + 1.35(C_meas-C0)] / 4，其中 q_base = 0.5。")

    doc.add_heading("5 仿真参数与运行环境", level=1)
    add_table(doc, ["参数", "数值"], [
        ["仿真时间", "2400 s"],
        ["采样步长", "0.5 s"],
        ["扰动加入时刻", "1200 s"],
        ["供气流量对象时间常数", "15 s"],
        ["DO 对象时间常数", "180 s"],
        ["纯滞后时间", "30 s"],
        ["鼓风机指令范围", "0 到 1"],
    ])
    add_body_paragraph(doc, "本报告对应的仿真程序包括 Python 版本和 MATLAB 版本。Python 脚本为 python/sim_method3_feedforward_cascade.py，MATLAB 脚本为 matlab/sim_method3_feedforward_cascade.m，二者采用相同模型和参数。")

    doc.add_heading("6 仿真结果与图像分析", level=1)
    add_picture_if_exists(doc, PY_FIG / "method3_fig1_do_disturbance_response.png", "图2 进水负荷阶跃扰动下的 DO 响应", width=6.3)
    add_body_paragraph(doc, "图2是方法三的核心结果。进水负荷突增后，普通串级 PID 的 DO 下跌较深且恢复较慢；普通前馈 + 串级 PID 能提前提高供气量，因此 DO 最低值更高，恢复到设定值附近的时间更短。")

    add_picture_if_exists(doc, PY_FIG / "method3_fig2_air_setpoint.png", "图3 目标供气量信号对比", width=6.3)
    add_body_paragraph(doc, "图3表明，普通串级 PID 的目标供气量主要由 DO 偏差反馈产生，因此变化滞后。加入前馈后，系统在检测到进水流量和污染物浓度升高时，直接提高目标供气量，实现提前补偿。")

    add_picture_if_exists(doc, PY_FIG / "method3_fig3_feedforward_signals.png", "图4 前馈输入与前馈输出信号", width=6.2)
    add_body_paragraph(doc, "图4展示了前馈补偿的来源。进水流量和污染物浓度阶跃上升后，前馈模型计算得到的供气量同步上升，说明控制器将耗氧端扰动转换成了供气端补偿。")

    add_picture_if_exists(doc, PY_FIG / "method3_fig4_control_signal.png", "图5 鼓风机控制指令对比", width=6.3)
    add_body_paragraph(doc, "图5说明前馈 + 串级 PID 在扰动初期控制指令更快上升，鼓风机提前加大输出。这种策略用短时间内略高的供气量换取更小的 DO 偏差和更快的恢复速度。")

    add_picture_if_exists(doc, PY_FIG / "method3_fig5_cascade_inner_loop.png", "图6 供气流量内环跟踪效果", width=6.3)
    add_body_paragraph(doc, "图6反映供气流量内环的跟踪能力。前馈通道提高目标供气量后，内环 PID 快速调节鼓风机或阀门，使实际供气量跟随目标值变化。")

    add_picture_if_exists(doc, PY_FIG / "method3_fig6_load_disturbance.png", "图7 进水耗氧负荷扰动及来源", width=6.3)
    add_body_paragraph(doc, "图7给出了扰动工况。进水流量和污染物浓度同时升高，导致耗氧负荷增加。上半图中的负向扰动表示该扰动会消耗更多氧气，使 DO 产生下降趋势。")

    add_picture_if_exists(doc, PY_FIG / "method3_fig7_performance_bar.png", "图8 抗扰动性能指标对比", width=6.3)
    add_body_paragraph(doc, "图8将最大偏差、恢复时间、IAE 和控制能量进行对比。前馈 + 串级 PID 的最大偏差、恢复时间和 IAE 均低于普通串级 PID，说明其抗扰动能力更强；控制能量略高，体现了提前供氧带来的能耗代价。")

    doc.add_heading("7 定量指标对比", level=1)
    add_table(doc, ["控制方法", "最低 DO / mg/L", "最大偏差 / mg/L", "恢复时间 / s", "IAE", "控制能量"], [
        ["串级 PID", "1.783", "0.217", "448.5", "66.79", "613.42"],
        ["普通前馈 + 串级 PID", "1.818", "0.182", "240.0", "36.48", "628.32"],
    ])
    add_body_paragraph(doc, "从定量指标看，普通前馈 + 串级 PID 将最大 DO 偏差由 0.217 mg/L 降低至 0.182 mg/L，将恢复时间由 448.5 s 缩短至 240.0 s，IAE 由 66.79 降低至 36.48。说明该方法能够显著改善进水负荷扰动下的动态性能。")

    doc.add_heading("8 结论", level=1)
    add_body_paragraph(doc, "方法三普通前馈 + 串级 PID 在串级控制基础上引入进水侧扰动补偿，能够在 DO 明显下降前提前增加目标供气量，使供气流量内环提前动作。与单纯串级 PID 相比，该方法具有更小的 DO 下跌幅度、更短的恢复时间和更低的累积误差。")
    add_body_paragraph(doc, "由于前馈模型难以完全准确，且在线测量存在滤波和滞后，前馈控制不能单独替代反馈控制。因此，保留 DO 外环反馈是必要的。总体来看，普通前馈 + 串级 PID 适合用于进水流量和污染物浓度可以在线测量或估算的污水处理曝气系统。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
